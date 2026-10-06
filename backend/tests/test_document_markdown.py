"""Document bodies are markdown: docx / pdf / pptx get real structure."""

import io

import pytest

from app.services.document_markdown import parse_blocks, parse_inline, plain_text
from app.services.document_render import render_document

BODY = (
    "Intro line.\n\n"
    "# Overview\n\n"
    "Revenue grew **12%** and *costs* held.\n\n"
    "- supplier delays\n- staffing\n\n"
    "## Next steps\n\n"
    "1. finalize plan\n2. review\n"
)


def test_parse_blocks_kinds():
    kinds = [b.kind for b in parse_blocks(BODY)]
    assert kinds == [
        "paragraph",
        "heading",
        "paragraph",
        "bullet",
        "bullet",
        "heading",
        "number",
        "number",
    ]


def test_parse_inline_styles_and_links():
    runs = parse_inline("a **b** *c* `d` [e](http://x)")
    assert any(r.text == "b" and r.bold for r in runs)
    assert any(r.text == "c" and r.italic for r in runs)
    assert any(r.text == "d" and r.code for r in runs)
    link = [r for r in runs if r.url]
    assert link and link[0].text == "e" and link[0].url == "http://x"
    assert "http" not in "".join(r.text for r in runs)


def test_plain_body_is_still_paragraphs():
    blocks = parse_blocks("one\n\ntwo")
    assert [b.text for b in blocks] == ["one", "two"]


def test_docx_has_headings_lists_and_bold():
    from docx import Document

    data = render_document("T", BODY, "docx")
    doc = Document(io.BytesIO(data))
    styles = [p.style.name for p in doc.paragraphs]
    assert "Heading 1" in styles and "Heading 2" in styles
    assert styles.count("List Bullet") == 2
    # numbered items restart under each heading and are written out
    assert [p.text for p in doc.paragraphs if p.text[:3] in ("1. ", "2. ")] == [
        "1.  finalize plan",
        "2.  review",
    ]
    assert not any("**" in p.text or "# " in p.text for p in doc.paragraphs)
    bold = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "12%" in bold


def test_pptx_headings_become_slides():
    from pptx import Presentation

    prs = Presentation(io.BytesIO(render_document("T", BODY, "pptx")))
    titles = [s.shapes.title.text for s in prs.slides]
    # BODY uses "#" for the section and "##" inside it: one slide, sub-heading kept.
    assert titles == ["T", "Overview"]


def test_pdf_builds():
    pytest.importorskip("reportlab")
    data = render_document("T", BODY, "pdf")
    assert data.startswith(b"%PDF")


def test_markdown_passthrough():
    assert "# Overview" in render_document("T", BODY, "markdown").decode()


def test_editor_blank_paragraphs_and_entities_do_not_leak():
    blocks = parse_blocks("Title\n\n&nbsp;\n\nA &amp; B\n\n&#160;\n\nEnd")
    assert [b.text for b in blocks] == ["Title", "A & B", "End"]
    runs = parse_inline("5&nbsp;kg")
    assert "&nbsp;" not in "".join(r.text for r in runs)


RICH = (
    "Summary line.\n\n"
    "# Overview\n\n"
    "> Key point: costs held.\n\n"
    "| Item | Cost |\n|---|---|\n| Hardware | 40k |\n| Staffing | 120k |\n\n"
    "- top\n  - nested\n\n"
    "Visit [the site](https://example.com) and run `make`.\n\n"
    "```\ncode block\n```\n"
)


def test_parse_rich_blocks():
    kinds = [b.kind for b in parse_blocks(RICH)]
    assert kinds == [
        "paragraph",
        "heading",
        "quote",
        "table",
        "bullet",
        "bullet",
        "paragraph",
        "code",
    ]
    table = [b for b in parse_blocks(RICH) if b.kind == "table"][0]
    assert len(table.rows) == 3 and table.rows[0][0][0].text == "Item"
    nested = [b for b in parse_blocks(RICH) if b.kind == "bullet"]
    assert [b.level for b in nested] == [0, 1]


def test_docx_is_styled_with_table_callout_link_and_footer():
    from docx import Document
    from docx.oxml.ns import qn

    data = render_document("Quarterly Plan", RICH, "docx")
    doc = Document(io.BytesIO(data))
    h1 = doc.styles["Heading 1"]
    assert str(h1.font.color.rgb) == "1F4E79" and h1.font.bold
    assert doc.styles["Title"].font.size.pt == 28
    assert len(doc.tables) == 1
    header_fill = doc.tables[0].cell(0, 0)._tc.tcPr.find(qn("w:shd")).get(qn("w:fill"))
    assert header_fill == "1F4E79"
    quote = [p for p in doc.paragraphs if p.text.startswith("Key point")][0]
    assert (
        quote._p.pPr.find(qn("w:shd")) is not None
        and quote._p.pPr.find(qn("w:pBdr")) is not None
    )
    assert doc.element.xpath("//w:hyperlink")
    assert "Page" in doc.sections[0].footer.paragraphs[0].text


def test_theme_accent_is_configurable_and_invalid_falls_back():
    from docx import Document

    doc = Document(
        io.BytesIO(
            render_document("T", "# A", "docx", theme={"accent_color": "#0B6E4F"})
        )
    )
    assert str(doc.styles["Heading 1"].font.color.rgb) == "0B6E4F"
    doc = Document(
        io.BytesIO(
            render_document("T", "# A", "docx", theme={"accent_color": "not-a-colour"})
        )
    )
    assert str(doc.styles["Heading 1"].font.color.rgb) == "1F4E79"


def _letterhead():
    from docx import Document

    d = Document()
    d.sections[0].header.paragraphs[0].text = "ACME CORP - 1 Main Street"
    d.add_paragraph("sample text that must not survive")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_letterhead_template_keeps_header_and_drops_sample_body():
    from docx import Document

    data = render_document("T", "# Scope\n\nbody", "docx", template=_letterhead())
    doc = Document(io.BytesIO(data))
    assert doc.sections[0].header.paragraphs[0].text == "ACME CORP - 1 Main Street"
    assert not any("sample text" in p.text for p in doc.paragraphs)
    assert any(p.text == "Scope" for p in doc.paragraphs)


def test_wrong_kind_of_template_is_ignored():
    from docx import Document

    pptx_like = render_document("T", "# A", "pptx")
    doc = Document(io.BytesIO(render_document("T", "# A", "docx", template=pptx_like)))
    assert any(p.text == "A" for p in doc.paragraphs)
    assert render_document("T", "# A", "docx", template=b"not a zip")


def test_pptx_is_widescreen_with_styled_title_slide_and_table():
    from pptx import Presentation

    prs = Presentation(
        io.BytesIO(
            render_document(
                "Deck", "Intro\n\n## Costs\n\n| A | B |\n|---|---|\n| 1 | 2 |", "pptx"
            )
        )
    )
    assert prs.slide_width > prs.slide_height * 1.7
    assert [s.shapes.title.text for s in prs.slides] == ["Deck", "Costs"]
    assert any(sh.has_table for sh in prs.slides[1].shapes)


def test_pdf_with_rich_content_builds():
    pytest.importorskip("reportlab")
    assert render_document("T", RICH, "pdf").startswith(b"%PDF")


def test_slides_start_at_the_shallowest_heading_level():
    from app.services.document_markdown import parse_blocks, split_into_slides

    only_h2 = parse_blocks("Intro\n\n## A\n\n- a1\n\n## B\n\n- b1")
    intro, slides = split_into_slides(only_h2)
    assert [t for t, _ in slides] == ["A", "B"] and len(intro) == 1

    nested = parse_blocks("# One\n\n## One-a\n\n- x\n\n## One-b\n\n- y\n\n# Two\n\n- z")
    _, slides = split_into_slides(nested)
    assert [t for t, _ in slides] == ["One", "Two"]
    assert [b.kind for b in slides[0][1]] == ["heading", "bullet", "heading", "bullet"]
    assert split_into_slides(parse_blocks("no headings at all"))[1] == []


def test_pptx_subheadings_stay_inside_their_slide():
    from pptx import Presentation

    body = "# Plan\n\n## Scope\n\n- a\n\n## Timeline\n\n- b"
    prs = Presentation(io.BytesIO(render_document("Deck", body, "pptx")))
    assert [s.shapes.title.text for s in prs.slides] == ["Deck", "Plan"]
    texts = " ".join(
        sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame
    )
    assert "Scope" in texts and "Timeline" in texts
