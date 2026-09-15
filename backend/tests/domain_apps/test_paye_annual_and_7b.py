"""GRA Form 2 (Annual PAYE Employer Return) + Form 7B (Annual Employee
Emolument Slip) — golden-shape, aggregation, and generation tests.

Mirrors test_paye_tools.py's ``_FakeEntry``/``_FakeCtx`` structure. Covers:

- The generated CSV/XLSX header rows match the verified column lists
  EXACTLY (byte-for-byte) — GRA's uploader matches headers literally, so
  this must never silently drift.
- The year-rollup aggregation sums RAW components across months and
  re-derives formula columns from the sums, rather than trusting (or
  summing) any already-computed value a monthly row happens to carry —
  proven by deliberately poisoning each monthly row's stored
  adjusted_7a/total_income/total_deductions with wrong values and
  confirming the aggregate ignores them.
- Form 2's employer summary row (year-only Other_Names, correct totals,
  correct employee count).
- Form 7B has NO employer summary row (unlike Form 2/5).
"""

from __future__ import annotations

import base64

import pytest

from app.profiles.payroll_app.tools import _7b_format as fmt7b  # type: ignore[import]
from app.profiles.payroll_app.tools import (
    _paye_annual_agg as agg,  # type: ignore[import]
)
from app.profiles.payroll_app.tools import _paye_format as fmt  # type: ignore[import]
from app.profiles.payroll_app.tools import generate_7b_template  # type: ignore[import]
from app.profiles.payroll_app.tools import paye_annual_generate  # type: ignore[import]

# ── golden shape ─────────────────────────────────────────────────────────

FORM2_VERIFIED_HEADERS = [
    "TIN",
    "Employee_Number",
    "First_Name",
    "Last_Name",
    "Other_Names",
    "Address",
    "Pay_Frequency",
    "Period_Employed",
    "Employee_Type",
    "Primary_Secondary_Job",
    "Value_7A_Salaries_Wages",
    "Total_Overtime",
    "Second_Job_Deduction",
    "Overtime_Deduction",
    "Adjusted_7A_Salaries_Wages",
    "Value_7B_Board_Lodge",
    "Value_7C_Other_Taxable_Allowances",
    "Value_7C_Other_Non_Taxable_Allowances",
    "Total_Income",
    "Personal_Allowance",
    "Employee_NIS_Contribution",
    "Medical_Life_Insurance_Premiums_Deduction",
    "Children_Deduction",
    "Total_Deductions",
    "Tax_Deducted",
    "Date_Of_Birth",
    "Bank_Name",
    "Bank_Account_No",
    "Bank_Account_Routing_Sort_code",
    "Bank_Account_Transit_No",
    "Child_Declaration_No",
]

SEVEN_B_VERIFIED_HEADERS = [
    "TIN",
    "Employee_Num",
    "First_Name",
    "Last_Name",
    "Other_Names",
    "Employee_address",
    "Primary_Secondary_Job",
    "Basic_salary",
    "Num_Overtime_Months",
    "Total_Overtime",
    "Second_Job_Deduction",
    "Overtime_Deduction",
    "Adjusted_7A_Salaries_Wages",
    "Rentfree_quarters or House allowances",
    "Board_Lodge",
    "Fees",
    "Bonus and Profit_share(incentives)",
    "All_Other Taxable_allowances",
    "Total_Taxable_Allowances",
    "Total_taxable_income",
    "Total_Non_Taxable_Allowances",
    "Total_Income",
    "Employee_NIS_Contribution",
    "Medical_Life_Insurance_Premiums_Deduction",
    "Num_Children",
    "Children_Deduction",
    "Total_Deductions",
    "Tax_Deducted",
    "National_ID",
    "Period_from(yyyy-mm-dd)",
    "Period_To(yyyy-mm-dd)",
    "Employer_TIN",
    "Company_Name",
    "Company_address",
    "Employee_email",
]


def test_form2_reuses_form5s_identical_31_column_schema():
    """Form 2 and Form 5 share ONE column schema (_paye_format.COLUMN_HEADERS)
    — this is the verified byte-for-byte GRA header row, transcribed from
    the real template file given for this task."""
    assert fmt.COLUMN_HEADERS == FORM2_VERIFIED_HEADERS


def test_7b_column_headers_match_real_template_byte_for_byte():
    """Verified against the real 7B Generator Template.xlsx's row 1 via
    openpyxl (see _7b_format.py's module docstring) — including the odd
    spacing/casing GRA's own tool matches literally."""
    assert fmt7b.COLUMN_HEADERS == SEVEN_B_VERIFIED_HEADERS
    assert len(fmt7b.FIELD_KEY_ORDER) == len(fmt7b.COLUMN_HEADERS) == 35


# ── Form 2 aggregation ───────────────────────────────────────────────────


def _monthly_paye_line(**overrides):
    base = {
        "employee": "emp-1",
        "tin": "987654321",
        "employee_number": "E1",
        "first_name": "Jane",
        "last_name": "Doe",
        "other_names": "",
        "address": "2 Second St",
        "pay_frequency": "Monthly",
        "period_employed": "1",
        "employee_type": "Full-Time",
        "primary_secondary_job": "Primary",
        "value_7a": 100000,
        "total_overtime": 5000,
        "second_job_deduction": 1000,
        "overtime_deduction": 500,
        "value_7b": 2000,
        "value_7c_taxable": 1000,
        "value_7c_nontaxable": 500,
        "personal_allowance": 65000,
        "employee_nis_contribution": 5600,
        "medical_life_insurance": 1000,
        "children_deduction": 2000,
        "tax_deducted": 3000,
        "date_of_birth": "1990-01-01",
        "bank_name": "Bank Co",
        "account_no": "1111",
        "routing_transit": "2222",
        "child_declaration_no": "",
        # Deliberately WRONG stored formula fields — a real monthly row
        # would have these correctly set by paye_calc_row, but the
        # aggregator must never trust (or sum) them; it must recompute
        # from the raw fields it itself sums.
        "adjusted_7a": -1,
        "total_income": -1,
        "total_deductions": -1,
    }
    base.update(overrides)
    return base


def test_aggregate_form2_employee_sums_raw_fields_and_recomputes_formulas():
    jan = _monthly_paye_line(period_employed="1")
    feb = _monthly_paye_line(period_employed="1", tax_deducted=3200)

    rows = agg.aggregate_form2_year([jan, feb])

    assert len(rows) == 1
    row = rows[0]

    # Raw components summed across the two months.
    assert row["value_7a"] == 200000
    assert row["total_overtime"] == 10000
    assert row["second_job_deduction"] == 2000
    assert row["overtime_deduction"] == 1000
    assert row["value_7b"] == 4000
    assert row["personal_allowance"] == 130000
    assert row["employee_nis_contribution"] == 11200
    assert row["tax_deducted"] == 3000 + 3200
    # Period_Employed summed too (both months parse as plain numbers).
    assert row["period_employed"] == 2

    # Formula columns re-derived from the SUMMED raw fields — NOT the
    # poisoned -1 values each monthly row carried, and NOT a naive sum of
    # those poisoned values either (which would be -2).
    expected_adjusted_7a = 200000 + 10000 - 2000 - 1000
    assert row["adjusted_7a"] == expected_adjusted_7a
    expected_total_income = expected_adjusted_7a + 4000 + 2000 + 1000
    assert row["total_income"] == expected_total_income
    expected_total_deductions = 130000 + 11200 + 2000 + 4000
    assert row["total_deductions"] == expected_total_deductions

    # Identity fields carried forward (latest month wins — here both
    # identical).
    assert row["first_name"] == "Jane"
    assert row["last_name"] == "Doe"
    assert row["tin"] == "987654321"


def test_aggregate_form2_year_groups_by_employee_not_by_row():
    emp1_jan = _monthly_paye_line(tin="111", employee_number="E1", first_name="Jane")
    emp1_feb = _monthly_paye_line(tin="111", employee_number="E1", first_name="Jane")
    emp2_jan = _monthly_paye_line(tin="222", employee_number="E2", first_name="John")

    rows = agg.aggregate_form2_year([emp1_jan, emp1_feb, emp2_jan])

    assert len(rows) == 2
    by_tin = {r["tin"]: r for r in rows}
    assert by_tin["111"]["value_7a"] == 200000  # two months
    assert by_tin["222"]["value_7a"] == 100000  # one month


def test_form2_employer_totals_sum_the_aggregated_rows():
    jan = _monthly_paye_line(tin="111")
    feb = _monthly_paye_line(tin="222")
    rows = agg.aggregate_form2_year([jan, feb])
    totals = agg.form2_employer_totals(rows)
    assert totals["value_7a"] == 100000 + 100000
    assert totals["adjusted_7a"] == sum(r["adjusted_7a"] for r in rows)


# ── Form 2 employer summary row (via _paye_format.generate_csv annual=True) ─


def test_generate_csv_annual_summary_row_uses_year_only_other_names():
    row = _monthly_paye_line(tax_deducted=3000)
    del row["adjusted_7a"], row["total_income"], row["total_deductions"]

    csv_text = fmt.generate_csv(
        company_name="Acme Co",
        company_tin="123456789",
        company_address="1 Main St",
        year=2026,
        period=None,
        lines=[row],
        annual=True,
    )
    rows = csv_text.strip("\n").split("\n")
    assert rows[0] == ",".join(fmt.COLUMN_HEADERS)
    summary = rows[2].split(",")
    assert summary[0] == "123456789"  # company TIN
    assert summary[2] == "Acme Co"  # company name
    # Year only — NOT "'06/2026" the way the monthly form renders it.
    assert summary[4] == "2026"
    assert summary[7] == "1"  # employee count


def test_build_filename_annual_matches_gra_pattern():
    assert fmt.build_filename_annual("Acme Co", 2026) == "paye-AcmeCo-2026.csv"


# ── Form 7B aggregation ──────────────────────────────────────────────────


def _monthly_7b_row(**overrides):
    base = agg.paye_filing_line_to_7b_row(_monthly_paye_line())
    # Poison the formula columns the same way, to prove re-derivation.
    base["adjusted_7a"] = -1
    base["total_taxable_allowances"] = -1
    base["total_taxable_income"] = -1
    base["total_income"] = -1
    base["total_deductions"] = -1
    base.update(overrides)
    return base


def test_aggregate_7b_employee_sums_and_recomputes_formulas():
    jan = _monthly_7b_row()
    feb = _monthly_7b_row()

    rows = agg.aggregate_7b_year([jan, feb])

    assert len(rows) == 1
    row = rows[0]
    assert row["basic_salary"] == 200000
    assert row["total_overtime"] == 10000

    expected_adjusted_7a = 200000 + 10000 - 2000 - 1000
    assert row["adjusted_7a"] == expected_adjusted_7a
    # value_7b -> board_lodge; value_7c_taxable -> other_taxable_allowances
    expected_taxable_allowances = 0 + (2000 * 2) + 0 + 0 + (1000 * 2)
    assert row["total_taxable_allowances"] == expected_taxable_allowances
    expected_taxable_income = expected_adjusted_7a + expected_taxable_allowances
    assert row["total_taxable_income"] == expected_taxable_income
    expected_non_taxable = 500 * 2
    assert row["total_non_taxable_allowances"] == expected_non_taxable
    assert row["total_income"] == expected_taxable_income + expected_non_taxable
    expected_deductions = (5600 * 2) + (1000 * 2) + (2000 * 2)
    assert row["total_deductions"] == expected_deductions


def test_aggregate_7b_employee_resolves_national_id_and_email_from_hr_employee():
    jan = _monthly_7b_row()
    employee_cf_by_id = {
        "emp-1": {"id_number": "NID-123", "personal_email": "jane@example.com"}
    }
    rows = agg.aggregate_7b_year([jan], employee_cf_by_id=employee_cf_by_id)
    assert rows[0]["national_id"] == "NID-123"
    assert rows[0]["employee_email"] == "jane@example.com"


def test_7b_generate_csv_has_no_employer_summary_row():
    jan = _monthly_7b_row()
    feb_other_employee = _monthly_7b_row()
    feb_other_employee["tin"] = "222"
    feb_other_employee["employee_num"] = "E2"
    rows = agg.aggregate_7b_year([jan, feb_other_employee])

    csv_text = fmt7b.generate_csv(rows)
    lines = csv_text.strip("\n").split("\n")
    assert lines[0] == ",".join(fmt7b.COLUMN_HEADERS)
    # Exactly one data row per employee, no trailing summary row.
    assert len(lines) == 1 + len(rows) == 3


def test_7b_generate_xlsx_bytes_has_correct_header_row():
    from io import BytesIO

    import openpyxl

    rows = agg.aggregate_7b_year([_monthly_7b_row()])
    xlsx_bytes = fmt7b.generate_xlsx_bytes(rows)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header == fmt7b.COLUMN_HEADERS
    assert ws.max_row == 2  # header + one employee row, no summary row


# ── Full tool round trip (aggregate_from_monthly_filings + generate) ──────


class _FakeEntry:
    def __init__(self, entry_id, custom_fields, track_id="track-1"):
        self.id = entry_id
        self.custom_fields = custom_fields
        self.track_id = track_id


class _FakeCtx:
    def __init__(self, *, filings, lines, employees=None):
        self._filings = list(filings)
        self._lines = list(lines)
        self._employees = list(employees or [])
        self.updates: dict[str, dict] = {}
        self.created: list[_FakeEntry] = []

    async def find_entries_in_track_type(self, track_type):
        return {"filings": self._filings}.get(track_type, [])

    async def find_entries(self, query):
        track_id = query.get("track_id")
        return [e for e in self._lines if e.track_id == track_id]

    async def get_entry_system(self, entry_id):
        for e in [*self._filings, *self._lines, *self._employees]:
            if e.id == entry_id:
                return e
        return None

    async def update_entry_fields(self, entry_id, custom_fields):
        self.updates.setdefault(entry_id, {}).update(custom_fields)
        for e in [*self._filings, *self._lines]:
            if e.id == entry_id:
                e.custom_fields = {**(e.custom_fields or {}), **custom_fields}
        return True

    async def create_entry(self, *, track_id, entry_type_key, title, custom_fields):
        new_id = f"created-{len(self.created)}"
        entry = _FakeEntry(new_id, dict(custom_fields), track_id=track_id)
        self.created.append(entry)
        self._lines.append(entry)
        return entry


@pytest.mark.asyncio
async def test_paye_annual_aggregate_and_generate_csv_round_trip():
    jan_filing = _FakeEntry(
        "filing-jan",
        {
            "company_tin": "123456789",
            "period": 1,
            "year": 2026,
            "employee_lines_track": "track-jan",
        },
    )
    feb_filing = _FakeEntry(
        "filing-feb",
        {
            "company_tin": "123456789",
            "period": 2,
            "year": 2026,
            "employee_lines_track": "track-feb",
        },
    )
    jan_line = _FakeEntry("line-jan-1", _monthly_paye_line(), track_id="track-jan")
    feb_line = _FakeEntry("line-feb-1", _monthly_paye_line(), track_id="track-feb")

    annual_return = _FakeEntry(
        "annual-1",
        {
            "company_name": "Acme Co",
            "company_tin": "123456789",
            "company_address": "1 Main St",
            "year": 2026,
            "employee_lines_track": "track-annual",
            "status": "draft",
        },
    )

    ctx = _FakeCtx(
        filings=[jan_filing, feb_filing, annual_return], lines=[jan_line, feb_line]
    )

    agg_result = await paye_annual_generate.aggregate_from_monthly_filings(
        {"entry_id": "annual-1"}, ctx
    )
    assert agg_result["ok"] is True
    assert agg_result["aggregated_count"] == 1  # one employee across both months
    assert agg_result["created_count"] == 1
    assert agg_result["source_months"] == 2

    gen_result = await paye_annual_generate.generate_csv({"entry_id": "annual-1"}, ctx)
    assert gen_result["ok"] is True
    assert gen_result["file"]["filename"] == "paye-AcmeCo-2026.csv"
    csv_text = base64.b64decode(gen_result["file"]["content_base64"]).decode("utf-8")
    rows = csv_text.strip("\n").split("\n")
    assert rows[0] == ",".join(fmt.COLUMN_HEADERS)
    assert len(rows) == 3  # header + 1 employee + 1 summary
    data_row = rows[1].split(",")
    assert data_row[10] == "200000"  # Value_7A summed across Jan+Feb
    summary_row = rows[2].split(",")
    assert summary_row[4] == "2026"  # year-only Other_Names
    assert annual_return.custom_fields["status"] == "generated"

    # Re-running the aggregate is safe — updates in place, doesn't duplicate.
    agg_result_2 = await paye_annual_generate.aggregate_from_monthly_filings(
        {"entry_id": "annual-1"}, ctx
    )
    assert agg_result_2["created_count"] == 0
    assert agg_result_2["updated_count"] == 1
    assert len(ctx.created) == 1


@pytest.mark.asyncio
async def test_paye_7b_aggregate_and_generate_xlsx_round_trip():
    from io import BytesIO

    import openpyxl

    jan_filing = _FakeEntry(
        "filing-jan",
        {
            "company_tin": "123456789",
            "period": 1,
            "year": 2026,
            "employee_lines_track": "track-jan",
        },
    )
    jan_line = _FakeEntry("line-jan-1", _monthly_paye_line(), track_id="track-jan")

    batch = _FakeEntry(
        "batch-1",
        {
            "company_name": "Acme Co",
            "company_tin": "123456789",
            "company_address": "1 Main St",
            "year": 2026,
            "employee_lines_track": "track-7b",
            "status": "draft",
        },
    )
    employee = _FakeEntry(
        "emp-1", {"id_number": "NID-1", "personal_email": "jane@example.com"}
    )

    ctx = _FakeCtx(filings=[jan_filing, batch], lines=[jan_line], employees=[employee])

    agg_result = await generate_7b_template.aggregate_from_monthly_filings(
        {"entry_id": "batch-1"}, ctx
    )
    assert agg_result["ok"] is True
    assert agg_result["aggregated_count"] == 1
    assert agg_result["created_count"] == 1

    gen_result = await generate_7b_template.generate_xlsx({"entry_id": "batch-1"}, ctx)
    assert gen_result["ok"] is True
    assert gen_result["file"]["filename"] == "7b-AcmeCo-2026.xlsx"
    xlsx_bytes = base64.b64decode(gen_result["file"]["content_base64"])
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header == fmt7b.COLUMN_HEADERS
    assert ws.max_row == 2  # header + one employee, no summary row
    data_row = [c.value for c in ws[2]]
    national_id_idx = fmt7b.FIELD_KEY_ORDER.index("national_id")
    email_idx = fmt7b.FIELD_KEY_ORDER.index("employee_email")
    assert data_row[national_id_idx] == "NID-1"
    assert data_row[email_idx] == "jane@example.com"
    assert batch.custom_fields["status"] == "generated"
