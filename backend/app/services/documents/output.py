"""Output renderers: HTML → PDF / DOCX / raw HTML bytes."""

from __future__ import annotations

import hashlib
import io
import logging
import re
from typing import Tuple

logger = logging.getLogger(__name__)


def _html_entities_to_text(text: str) -> str:
    return (
        (text or "")
        .replace("&nbsp;", " ")
        .replace("&amp;", "&")
        .replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&quot;", '"')
        .replace("&#39;", "'")
    )


def _html_to_plaintext(html: str) -> str:
    """Convert merged document HTML to plain text for ReportLab/DOCX fallback."""
    text = html or ""
    body_match = re.search(r"<body\b[^>]*>(.*)</body>", text, flags=re.I | re.DOTALL)
    if body_match:
        text = body_match.group(1)
    text = re.sub(
        r"<(script|style|head)\b[^>]*>.*?</\1>",
        "",
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
    text = re.sub(r"</(p|div|h[1-6]|li|tr|blockquote)>", "\n", text, flags=re.I)
    text = re.sub(r"<h[1-6]\b[^>]*>", "\n\n", text, flags=re.I)
    text = re.sub(r"<[^>]+>", "", text)
    text = _html_entities_to_text(text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _strip_tags(html: str) -> str:
    """Backward-compatible alias — prefer ``_html_to_plaintext``."""
    return _html_to_plaintext(html)


def render_html_bytes(html: str) -> bytes:
    return (html or "").encode("utf-8")


def _normalize_reportlab_markup(fragment: str) -> str:
    """Map sanitized HTML inline tags to ReportLab Paragraph markup."""
    text = fragment or ""

    def _swap_tag(name: str, repl: str, raw: str) -> str:
        out = re.sub(
            rf"</?{name}\b[^>]*>",
            lambda m: f"</{repl}>" if m.group(0)[1] == "/" else f"<{repl}>",
            raw,
            flags=re.I,
        )
        return out

    text = _swap_tag("strong", "b", text)
    text = _swap_tag("b", "b", text)
    text = _swap_tag("em", "i", text)
    text = _swap_tag("i", "i", text)
    text = _swap_tag("u", "u", text)

    def _bold_span(m: re.Match[str]) -> str:
        return f"<b>{m.group(1)}</b>"

    text = re.sub(
        r'<span[^>]*style="[^"]*font-weight:\s*(?:700|bold)[^"]*"[^>]*>(.*?)</span>',
        _bold_span,
        text,
        flags=re.I | re.DOTALL,
    )
    text = re.sub(r"<span[^>]*>", "", text, flags=re.I)
    text = re.sub(r"</span>", "", text, flags=re.I)
    return text.strip()


def _paragraph_alignment(tag_html: str) -> int:
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT

    m = re.search(r"text-align:\s*(left|right|center|justify)", tag_html, flags=re.I)
    if not m:
        return TA_LEFT
    align = m.group(1).lower()
    if align == "center":
        return TA_CENTER
    if align == "right":
        return TA_RIGHT
    if align == "justify":
        return TA_JUSTIFY
    return TA_LEFT


def _signature_fragments_from_html(fragment: str) -> Tuple[str, str]:
    """Extract underline line + label from a ``doc-signature`` HTML fragment."""
    line_m = re.search(
        r'class="doc-signature-line"[^>]*>([^<]*)',
        fragment or "",
        flags=re.I,
    )
    label_m = re.search(
        r'class="doc-signature-label"[^>]*>([^<]*)',
        fragment or "",
        flags=re.I,
    )
    line = _html_entities_to_text((line_m.group(1) if line_m else "").strip())
    label = _html_entities_to_text((label_m.group(1) if label_m else "").strip())
    return line, label


def _reportlab_signature_flowables(fragment: str, styles, *, index: int) -> list:
    """Render signature placeholder blocks for ReportLab fallback PDFs."""
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Spacer

    line, label = _signature_fragments_from_html(fragment)
    if not line and not label:
        return []

    out = []
    if line:
        out.append(Paragraph(line.replace("&", "&amp;"), styles["BodyText"]))
        out.append(Spacer(1, 4))
    if label:
        label_style = ParagraphStyle(
            name=f"DocSigLabel{index}",
            parent=styles["BodyText"],
            fontSize=9,
            textColor=colors.HexColor("#555555"),
        )
        out.append(Paragraph(label.replace("&", "&amp;"), label_style))
    out.append(Spacer(1, 12))
    return out


def _reportlab_story_from_html(html: str, styles) -> list:
    """Build ReportLab flowables preserving basic block alignment and inline marks."""
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, Spacer

    body = html or ""
    body_match = re.search(r"<div class=['\"]doc-body['\"][^>]*>(.*)</div>", body, flags=re.I | re.DOTALL)
    if body_match:
        body = body_match.group(1)
    else:
        body_match = re.search(r"<body\b[^>]*>(.*)</body>", body, flags=re.I | re.DOTALL)
        if body_match:
            body = body_match.group(1)

    story = []
    block_re = re.compile(
        r"<(p|h[1-4]|blockquote|li)\b([^>]*)>(.*?)</\1>",
        flags=re.I | re.DOTALL,
    )
    # Must not match ``doc-signature-line`` / ``doc-signature-label`` inner divs.
    sig_open_re = re.compile(r'<div class="doc-signature"\s', flags=re.I)
    events: list[tuple[int, str, re.Match[str]]] = []
    for match in block_re.finditer(body):
        events.append((match.start(), "block", match))
    for match in sig_open_re.finditer(body):
        events.append((match.start(), "signature", match))
    events.sort(key=lambda item: item[0])

    sig_index = 0
    for _, kind, match in events:
        if kind == "block":
            tag = match.group(1)
            attrs = match.group(2) or ""
            inner = match.group(3) or ""
            inner = re.sub(r"<br\s*/?>", "<br/>", inner, flags=re.I)
            markup = _normalize_reportlab_markup(inner)
            if not markup:
                continue
            alignment = _paragraph_alignment(attrs)
            base = styles["BodyText"]
            if tag.lower().startswith("h"):
                base = styles["Heading2"]
            style = ParagraphStyle(
                name=f"DocBlock{len(story)}",
                parent=base,
                alignment=alignment,
            )
            story.append(Paragraph(markup, style))
            story.append(Spacer(1, 8))
            continue

        pos = match.start()
        snippet = body[pos : pos + 2500]
        story.extend(
            _reportlab_signature_flowables(snippet, styles, index=sig_index)
        )
        sig_index += 1

    if story:
        return story

    plain = _html_to_plaintext(html or "")
    for block in re.split(r"\n{2,}", plain):
        block = block.strip()
        if not block:
            continue
        safe = (
            block.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace("\n", "<br/>")
        )
        story.append(Paragraph(safe, styles["BodyText"]))
        story.append(Spacer(1, 8))
    return story


def render_pdf_bytes(html: str, *, title: str = "Document") -> bytes:
    """Render HTML to PDF — WeasyPrint when available, else ReportLab fallback."""
    try:
        from weasyprint import HTML

        return HTML(string=html or "").write_pdf()
    except ImportError:
        logger.warning(
            "WeasyPrint unavailable — using ReportLab fallback (install weasyprint "
            "and system libs for full template formatting)"
        )
    except Exception:  # noqa: BLE001
        logger.exception(
            "WeasyPrint PDF render failed — using ReportLab fallback "
            "(check cairo/pango on the server)"
        )

    try:
        from reportlab.lib.pagesizes import LETTER
    except ImportError as exc:
        raise RuntimeError(
            "PDF rendering is not configured on this server (install reportlab "
            "or weasyprint)"
        ) from exc
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=LETTER,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title=title,
    )
    styles = getSampleStyleSheet()
    story = _reportlab_story_from_html(html or "", styles)
    if not story:
        story = [Paragraph("(empty document)", styles["BodyText"])]
    doc.build(story)
    return buf.getvalue()


def render_docx_bytes(html: str, *, title: str = "Document") -> bytes:
    """Render a basic DOCX from HTML-derived plain text (Phase 8 starter)."""
    from docx import Document

    document = Document()
    document.core_properties.title = title
    plain = _strip_tags(html or "")
    for block in re.split(r"\n{2,}", plain):
        block = block.strip()
        if block:
            document.add_paragraph(block)
    if not plain.strip():
        document.add_paragraph("")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def render_output(
    html: str, output_format: str, *, title: str = "Document"
) -> Tuple[bytes, str, str]:
    """Return (bytes, mime, extension)."""
    fmt = (output_format or "pdf").lower()
    if fmt == "html":
        data = render_html_bytes(html)
        return data, "text/html; charset=utf-8", "html"
    if fmt == "docx":
        data = render_docx_bytes(html, title=title)
        return (
            data,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
        )
    data = render_pdf_bytes(html, title=title)
    return data, "application/pdf", "pdf"


def checksum_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()
